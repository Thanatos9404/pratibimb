"""Build Pratibimb_Presentation.pptx: one full-bleed image per slide (16:9).

    node docs/deck/render.mjs --png <dir>          # writes <dir>/slide-NN.jpg
    python docs/deck/make_pptx.py <dir>            # -> StackUnderflow_BIT-Mesra-Jaipur/Pratibimb_Presentation.pptx

Needs python-pptx. The PDF is the primary deliverable; this PPTX mirrors it for viewers that prefer PowerPoint.
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "StackUnderflow_BIT-Mesra-Jaipur" / "Pratibimb_Presentation.pptx"


def main(img_dir: str) -> None:
    images = sorted(Path(img_dir).glob("slide-*.jpg"))
    if not images:
        raise SystemExit(f"no slide-*.jpg in {img_dir}; run render.mjs with --png first")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 13.333 x 7.5 in (16:9)
    blank = prs.slide_layouts[6]
    for img in images:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(img), 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.core_properties.title = "Pratibimb: a hybrid physiology + ML digital twin for Type 2 Diabetes"
    prs.core_properties.author = "Team Stack Underflow (Yashvardhan Thanvi, Devesh Jain)"
    prs.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)} ({len(images)} slides, {OUT.stat().st_size / 1e6:.1f} MB)")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "slides")
